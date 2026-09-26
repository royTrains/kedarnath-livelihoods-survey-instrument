# -*- coding: utf-8 -*-
"""Question register: every question, bilingual, with the literature each one comes from.

Two outputs, same content:
  Question_register.xlsx  -- four sheets (questions, constructed variables, literature, modules)
  Question_register.pdf   -- the same, typeset

Built from dictionary.py and translations_hi.py like everything else, so it cannot drift from the
form that is actually in the field. This is the document to hand a supervisor, an ethics committee
or a referee who asks "where did this question come from" -- every asked item carries its source,
and the literature sheet inverts that so you can see what each source contributed.

The PDF is compiled with XELATEX, not pdflatex: the register prints the Hindi alongside the English
and pdflatex cannot set Devanagari. Nirmala UI ships with Windows and carries the script.
"""
import os, subprocess, sys, collections
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from dictionary import ROWS, LSETS, MODULES
from translations_hi import HI, HI_LSETS

XELATEX = "xelatex"
MODNAME = {c: t for c, t in MODULES}
asked = [r for r in ROWS if r["origin"] == "asked"]
constructed = [r for r in ROWS if r["origin"] == "constructed"]

TYPE = {"multi": "select multiple", "cat": "select one", "bin": "yes/no", "text": "free text",
        "money": "amount (Rs)", "count": "whole number", "num": "number", "date": "date", "id": "id"}


def opts(r, hindi=False):
    if not r["lset"]:
        return ""
    src = HI_LSETS.get(r["lset"], LSETS[r["lset"]]) if hindi else LSETS[r["lset"]]
    return " | ".join(f"{k} {v}" for k, v in src.items())


# ============================================================ EXCEL
wb = openpyxl.Workbook()
HEAD = Font(bold=True, color="FFFFFF", size=10)
FILL = PatternFill("solid", fgColor="1F3864")
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet(name, cols, rows, widths):
    ws = wb.create_sheet(name)
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, FILL, Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


sheet("Questions",
      ["#", "Module", "Module name", "Variable", "Type", "Question (English)", "Question (Hindi)",
       "Response options", "Asked only if", "Source / literature"],
      [[i, r["module"], MODNAME[r["module"]], r["name"], TYPE.get(r["kind"], r["kind"]),
        r["question"], HI.get(r["name"], ""), opts(r), r["skip"], r["source"]]
       for i, r in enumerate(asked, 1)],
      [5, 8, 26, 26, 14, 58, 58, 46, 40, 60])

sheet("Constructed variables",
      ["#", "Variable", "Label", "How it is built", "What it means", "Source / literature"],
      [[i, r["name"], r["label"], r["formula"], r["desc"], r["source"]]
       for i, r in enumerate(constructed, 1)],
      [5, 28, 44, 56, 78, 56])

lit = collections.defaultdict(list)
for r in asked:
    lit[r["source"]].append(r["name"])
sheet("Literature",
      ["Source / literature", "Items", "Variables drawing on it"],
      sorted(([k, len(v), ", ".join(v)] for k, v in lit.items()), key=lambda x: -x[1]),
      [78, 8, 96])

cnt = collections.Counter(r["module"] for r in asked)
sheet("Modules",
      ["Module", "Name", "Questions asked"],
      [[c, t, cnt.get(c, 0)] for c, t in MODULES] + [["", "TOTAL", len(asked)]],
      [10, 52, 16])

del wb["Sheet"]
xlsx = os.path.join(HERE, "Question_register.xlsx")
wb.save(xlsx)
print(xlsx)

# ============================================================ PDF (XeLaTeX, for the Devanagari)
def esc(s):
    s = str(s or "")
    for a, b in [("\\", "\\textbackslash{}"), ("&", "\\&"), ("%", "\\%"), ("$", "\\$"), ("#", "\\#"),
                 ("_", "\\_"), ("{", "\\{"), ("}", "\\}"), ("~", "\\textasciitilde{}"),
                 ("^", "\\textasciicircum{}")]:
        s = s.replace(a, b)
    return s


def hi(s):
    return "{\\hifont " + esc(s) + "}" if s else ""


tex = r"""\documentclass[9pt,a4paper,landscape]{article}
\usepackage[margin=1.4cm,landscape]{geometry}
\usepackage{fontspec,longtable,array,booktabs,xcolor}
\setmainfont{Latin Modern Roman}
\newfontfamily\hifont{Nirmala UI}
\definecolor{hdr}{HTML}{1F3864}
\setlength{\tabcolsep}{4pt}\renewcommand{\arraystretch}{1.18}
\usepackage{fancyhdr}\pagestyle{fancy}\fancyhf{}
\fancyhead[L]{\small Kedarnath Yatra Worker Survey --- Question Register}
\fancyhead[R]{\small\thepage}\renewcommand{\headrulewidth}{0.3pt}
\begin{document}
\begin{center}
{\LARGE\bfseries Kedarnath Yatra Worker Survey}\\[3pt]
{\large Question register, with sources}\\[8pt]
""" + f"{len(asked)} questions asked \\quad$\\cdot$\\quad {len(constructed)} variables constructed \\quad$\\cdot$\\quad {len(lit)} distinct sources" + r"""\\[3pt]
{\small Every question is shown in English and Hindi exactly as the tablet presents it. The source
column names the instrument or paper the item is taken from or adapted from.}
\end{center}
\vspace{4pt}
"""

for code, title in MODULES:
    rows = [r for r in asked if r["module"] == code]
    if not rows:
        continue
    tex += ("\\noindent{\\large\\bfseries\\color{hdr} Module " + esc(code) + " --- " + esc(title) +
            "}\\quad{\\small(" + str(len(rows)) + " questions)}\\\\[3pt]\n")
    tex += (r"\begin{longtable}{@{}p{0.5cm}p{3.3cm}p{1.7cm}p{7.3cm}p{5.6cm}p{5.4cm}@{}}" "\n"
            r"\toprule \footnotesize\textbf{\#} & \footnotesize\textbf{Variable} & "
            r"\footnotesize\textbf{Type} & \footnotesize\textbf{Question (English / Hindi)} & "
            r"\footnotesize\textbf{Options {\&} skip} & \footnotesize\textbf{Source}\\ \midrule"
            "\n\\endhead\n")
    for i, r in enumerate(asked, 1):
        if r["module"] != code:
            continue
        q = "\\footnotesize " + esc(r["question"])
        if HI.get(r["name"]):
            q += "\\par\\vspace{1pt}{\\footnotesize\\color{hdr}" + hi(HI[r["name"]]) + "}"
        o = "\\scriptsize " + (esc(opts(r)) if r["lset"] else "")
        if r["skip"]:
            o += ("\\par\\vspace{1pt}\\textit{Asked only if: " + esc(r["skip"]) + "}") if r["lset"] else \
                 ("\\textit{Asked only if: " + esc(r["skip"]) + "}")
        tex += (f"\\footnotesize {i} & \\footnotesize\\texttt{{" + esc(r["name"]) + "} & \\scriptsize "
                + esc(TYPE.get(r["kind"], r["kind"])) + " & " + q + " & " + o
                + " & \\scriptsize " + esc(r["source"]) + "\\\\\n")
    tex += "\\bottomrule\\end{longtable}\\vspace{6pt}\n"

tex += r"""\clearpage\noindent{\large\bfseries\color{hdr} Sources, and what each one contributed}\\[4pt]
\begin{longtable}{@{}p{9.5cm}p{1.2cm}p{12.5cm}@{}}
\toprule \footnotesize\textbf{Source} & \footnotesize\textbf{Items} & \footnotesize\textbf{Variables}\\ \midrule
\endhead
"""
for k, v in sorted(lit.items(), key=lambda x: -len(x[1])):
    tex += ("\\scriptsize " + esc(k) + " & \\scriptsize " + str(len(v)) + " & \\scriptsize\\texttt{"
            + esc(", ".join(v)) + "}\\\\\n")
tex += "\\bottomrule\\end{longtable}\n\\end{document}\n"

tp = os.path.join(HERE, "Question_register.tex")
open(tp, "w", encoding="utf-8").write(tex)
env = os.environ.copy()
env["PATH"] = os.pathsep.join(x for x in env.get("PATH", "").split(os.pathsep)
                              if not x.lower().rstrip("\\/").endswith(".exe"))
for _ in range(2):
    r = subprocess.run([XELATEX, "-interaction=nonstopmode", "-halt-on-error", "Question_register.tex"],
                       cwd=HERE, env=env, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if r.returncode != 0:
        print(r.stdout[-2500:])
        raise SystemExit("XeLaTeX failed")
for ext in (".aux", ".log", ".out"):
    p = os.path.join(HERE, "Question_register" + ext)
    if os.path.exists(p):
        os.remove(p)
print(os.path.join(HERE, "Question_register.pdf"))
print(f"{len(asked)} questions, {len(constructed)} constructed, {len(lit)} distinct sources")
