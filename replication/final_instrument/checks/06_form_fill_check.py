"""Fill the form 100 times and check the result is what the analyses need.

Two things are tested, and they are different questions:

  (A) SKIP-LOGIC CONFORMANCE. Reads the relevance expressions out of the built XLSForm itself --
      not out of dictionary.py's free-text `skip` field, which is prose and is not what Kobo runs --
      translates each into Python, and checks that every asked variable is present exactly when its
      own relevance rule says it should be. A variable that is present when it should have been
      skipped, or missing when it should have been asked, is a real form defect.

  (B) ANALYSIS READINESS. For each analysis the project actually runs, checks that every variable it
      needs exists in the built dataset and has enough non-missing values to estimate on. This is the
      "do our questions actually generate the worked example" test.

Run after 02_build_final_dataset.do. Exits non-zero if anything fails.
"""
import os, re, sys
import numpy as np
import pandas as pd
import openpyxl
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'questionnaire'))
from dictionary import LSETS, ROWS

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
N_DRAW = 100
SEED = 20260924

d = pd.read_stata(os.path.join(ROOT, "data", "kedarnath_final_n200_full.dta"), convert_categoricals=False)
raw = pd.read_csv(os.path.join(ROOT, "data", "raw_asked.csv"))
sample = raw.sample(n=min(N_DRAW, len(raw)), random_state=SEED).reset_index(drop=True)

fails, notes = [], []
def ok(msg):   print(f"  PASS  {msg}")
def bad(msg):  fails.append(msg); print(f"  FAIL  {msg}")
def note(msg): notes.append(msg); print(f"  note  {msg}")

# ---------------------------------------------------------------- (A) skip logic
print(f"\n(A) SKIP-LOGIC CONFORMANCE -- {len(sample)} simulated submissions, rules read from the XLSForm\n")

# Fields the dictionary marks "may be left blank" are genuinely optional -- a verbatim
# "other, specify" that the respondent could not answer. A blank there is correct data, not a
# skip-logic violation, so only the WRONGLY-PRESENT direction is checked for them.
from dictionary import ROWS as _ROWS
OPTIONAL = {r["name"] for r in _ROWS if "may be left blank" in r["skip"]}

ws = openpyxl.load_workbook(os.path.join(ROOT, "questionnaire", "Kedarnath_final_kobo.xlsx"))["survey"]
hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
ix = {h: i for i, h in enumerate(hdr)}
rules = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    name, rel, typ = row[ix["name"]], row[ix["relevant"]], row[ix["type"]]
    if name and rel and typ not in ("begin_group", "end_group", "calculate"):
        rules[name] = rel

def to_py(expr):
    """XLSForm relevance -> Python. Covers the forms actually used in this instrument.

    NOTE the parenthesisation step: pandas' & and | bind TIGHTER than the comparison operators, so a
    naive translation of "a = 1 or a = 2" becomes "a == 1 | a == 2", which Python reads as
    "a == (1 | a) == 2" and raises. Every atomic comparison has to be wrapped before and/or are
    substituted."""
    e = expr.replace("${", "S['").replace("}", "']")
    e = re.sub(r"(?<![<>!=])=(?!=)", "==", e)                       # single = is equality in XPath
    e = e.replace("!==", "!=").replace("<==", "<=").replace(">==", ">=")
    e = re.sub(r"S\['\w+'\]\s*(?:[<>!=]=|[<>])\s*-?\d+(?:\.\d+)?", lambda m: f"({m.group(0)})", e)
    e = e.replace(" and ", " & ").replace(" or ", " | ")
    # count-selected() on a select_multiple: non-empty export string means at least one box ticked
    e = re.sub(r"count-selected\(S\['(\w+)'\]\)\s*>\s*0",
               lambda m: "(S['" + m.group(1) + "'].astype(str).str.strip().replace('nan','') != '')", e)
    # selected(${multi}, 'code'): the export is a space-separated code string, so test membership of
    # the split tokens rather than a substring -- 'selected(x, "1")' must not match the code 10.
    e = re.sub(r"selected\(S\['(\w+)'\],\s*'(\w+)'\)",
               lambda m: ("S['" + m.group(1) + "'].astype(str).str.split().apply("
                          "lambda t: '" + m.group(2) + "' in t)"), e)
    return e

S = {c: sample[c] for c in sample.columns}
# calc_months_no_work is a hidden in-form calculate, not an exported column -- rebuild it
S["calc_months_no_work"] = sum((sample[f"status_m{m}"] == 8).astype(int) for m in range(1, 13))
S["calc_offseason_work"] = sum(sample[f"status_m{m}"].between(2, 7).astype(int) for m in range(1, 13))

checked = skipped_ok = 0
for name, rel in rules.items():
    if name not in sample.columns:
        note(f"{name}: in the form but not in raw_asked.csv (paradata or Kobo-only) -- not checked")
        continue
    try:
        should_ask = eval(to_py(rel))                   # noqa: S307 - our own generated expressions
    except Exception as ex:
        bad(f"{name}: could not evaluate relevance {rel!r} ({ex})")
        continue
    col = sample[name]
    present = col.notna() & (col.astype(str).str.strip() != "") if col.dtype == object else col.notna()
    wrong_present = int((present & ~should_ask).sum())
    wrong_absent  = 0 if name in OPTIONAL else int((~present & should_ask).sum())
    checked += 1
    if wrong_present or wrong_absent:
        bad(f"{name}: {wrong_present} answered when it should have been skipped, "
            f"{wrong_absent} skipped when it should have been asked  [rule: {rel}]")
    else:
        skipped_ok += 1
if skipped_ok == checked and checked:
    ok(f"all {checked} gated variables match their own XLSForm relevance rule on all {len(sample)} submissions")

# ---------------------------------------------------------------- (A2) rule ordering
# A question whose relevance names a variable asked LATER can never be shown: at the moment the form
# evaluates it, the gate is still unanswered. loan_collateral sat before loan_against_asset for
# exactly this reason and was silently unreachable -- no skip-logic check caught it, because the
# generator set both independently. Only a real end-to-end fill did.
print("")
print("(A2) RULE ORDERING -- is any gated question asked before its own gate?")
print("")
order = {}
for _r in ROWS:
    if _r["origin"] in ("asked", "paradata"):
        order.setdefault(_r["name"], len(order))
bad_order = 0
for name, rel in rules.items():
    if name not in order:
        continue
    for dep in set(re.findall(r"\$\{(\w+)\}", rel)):
        # the hidden calculates are not questions, but they READ the calendar, so a question gated
        # on one of them is really gated on status_m12 -- skipping them hid a real ordering bug
        if dep.startswith("calc_"):
            dep = "status_m12"
        if dep in order and order[dep] > order[name]:
            bad(f"{name} is gated on {dep}, which is asked LATER -- it can never appear")
            bad_order += 1
if not bad_order:
    ok(f"every gated question is asked after the variable that gates it ({len(rules)} rules)")

# ---------------------------------------------------------------- (B) analysis readiness
print("\n(B) ANALYSIS READINESS -- does the instrument generate what each analysis consumes?\n")

XVARS = ["education_years", "has_bank_account", "credit_institutional", "training_received", "smartphone_owned",
         "age", "hhsize", "employment_type", "years_in_yatra_work", "migrant",
         "shock_any", "yatra_income_share", "income_seasonality_cv", "health_access_tier"]
NEEDS = {
    "Monetary VEP (Chaudhuri 3-stage FGLS)":
        ["cons_pc_pm", "poor", "poor_sensitivity_cpi", "cons_pc_pm_narrow", "cons_pc_ae_pm"] + XVARS,
    "Multidimensional VEP (Alkire-Foster / Lyons et al.)":
        ["health_access_deprived", "n_health_insured", "child_school_dep", "education_years",
         "floor_material", "roof_material", "wall_material", "house_type", "electricity",
         "drinking_water", "water_on_premises", "water_deprived", "cooking_fuel",
         "mpi_mortality_dep", "mpi_maternal_dep", "mpi_schooling_dep", "mpi_attendance_dep",
         "mpi_housing_dep", "mpi_sanitation_dep", "mpi_electricity_dep", "mpi_bank_dep",
         "mpi_score", "mpi_poor", "toilet_type", "wall_material",
         "cooking_fuel_deprived", "mpi_asset_count", "mpi_asset_deprived", "rcsi_score",
         "food_coping_deprived"] + XVARS,
    "Employment VEP (Apablaza + Chaudhuri)":
        ["emp_dep_access", "emp_dep_comp", "emp_dep_sec", "emp_dep_stab", "emp_dep_cond",
         "emp_dep_count", "emp_dep_score", "emp_poor_k2", "job_situation",
         "job_permanence", "hours_week_yatra", "work_income_pm"] + XVARS,
    "Transferability / structural displacement":
        ["occupation", "occupation_detail", "prev_occ", "prev_occ_change", "target_occ", "tk_regular_n", "tk_prior_n",
         "best_alt_occupation", "task_cover_best", "task_retain_best", "skill_move_type",
         "n_other_activities", "n_other_act_checked", "trek_dependent", "gps_lat", "gps_lon"],
    "Shocks, coping and finance":
        ["shock_count", "shock_any", "shock_covariate", "shock_idiosyncratic",
         "coped_sold_assets", "coped_cut_consumption", "took_loan_12m", "n_health_insured",
         "n_life_insured", "home_rural_urban"],
    "Mobility and the closure regime":
        ["resp_returns_at_closure", "hh_at_home_place", "closure_labour_migrant", "stays_all_year",
         "split_household", "months_here", "months_home_base", "months_third_place",
         "months_away_total", "worked_away_in_closure", "worked_other_places",
         "would_move_for_work", "migrant", "cons_pc_denom_season", "cons_pc_onsite_pm"],
}
# variables that are legitimately missing for most rows because they are skip-gated, with the
# coverage each one should clear. child_school_dep is only defined for households that HAVE a child
# aged 6-14; Alkire-Foster treats a household with no eligible member as non-deprived, so partial
# coverage here is correct behaviour, not a defect.
GATED = {"prev_occ": 0.20, "target_occ": 0.50, "best_alt_occupation": 0.90, "task_cover_best": 0.90,
         "task_retain_best": 0.90, "skill_move_type": 0.90, "child_school_dep": 0.20}

for analysis, cols in NEEDS.items():
    missing = [c for c in cols if c not in d.columns]
    if missing:
        bad(f"{analysis}: variable(s) not built -> {missing}")
        continue
    thin = []
    for c in cols:
        cov = d[c].notna().mean()
        if cov < GATED.get(c, 0.95):
            thin.append(f"{c} {cov:.0%}")
    if thin:
        bad(f"{analysis}: present but too thin to estimate on -> {thin}")
    else:
        ok(f"{analysis}: all {len(cols)} required variables built and populated")

# ---------------------------------------------------------------- (C) new-variable behaviour
print("\n(C) DO THE NEW VARIABLES BEHAVE AS EXPECTED?\n")

# income fallback
fb = d.income_from_fallback == 1
if fb.sum() == 0:
    bad("income fallback: nobody took it -- the route is never exercised")
elif d.loc[fb, "income_m1"].notna().any():
    bad("income fallback: monthly earnings present for a fallback respondent")
elif d.income_seasonality_cv.isna().any():
    bad("income fallback: income_seasonality_cv missing for some rows (they would drop out of the FGLS)")
else:
    cv_fb, cv_cal = d.loc[fb, "income_seasonality_cv"].mean(), d.loc[~fb, "income_seasonality_cv"].mean()
    ok(f"income fallback: {fb.mean():.0%} took it; CV never missing; fallback CV {cv_fb:.2f} < calendar CV {cv_cal:.2f} as predicted"
       if cv_fb < cv_cal else f"income fallback works, but fallback CV {cv_fb:.2f} is NOT below calendar CV {cv_cal:.2f}")
    if cv_fb >= cv_cal:
        bad("fallback CV should be a lower bound on the calendar CV -- it is not")

# multi-select split
# derived from the dictionary, not hardcoded: the occupation list has already been renumbered
# once, and a hardcoded range silently under-counts rather than failing loudly.
act_cols = [f"other_act_{k}" for k in sorted(LSETS["occ"])]
if not all(c in d.columns for c in act_cols):
    bad("multi-select: other_act_1..13 not all built")
else:
    recomputed = d[act_cols].sum(1)
    if not (recomputed == d.n_other_act_checked).all():
        bad("multi-select: n_other_act_checked does not equal the sum of the binaries")
    elif (d.loc[d.n_other_activities == 0, act_cols].sum(1) != 0).any():
        bad("multi-select: a respondent with 0 other activities has a ticked box")
    else:
        ok(f"multi-select: splits cleanly; {(d.n_other_activities>0).mean():.0%} hold >=1 other activity, "
           f"max {int(d.n_other_activities.max())}, stated-vs-ticked mismatch {d.other_act_mismatch.mean():.1%}")

# health_access_tier from GPS
lat_by_tier = d.groupby("health_access_tier")["gps_lat"].agg(["min", "max"])
if d.health_access_tier.isna().any():
    bad("health_access_tier: missing for some rows after dropping location_cluster")
elif not (lat_by_tier.loc[3, "max"] < 30.58 <= lat_by_tier.loc[2, "min"]):
    bad(f"health_access_tier: GPS bands overlap\n{lat_by_tier}")
else:
    ok(f"health_access_tier: derived from gps_lat with clean bands, shares {d.health_access_tier.value_counts(normalize=True).sort_index().round(2).to_dict()}")

# skill-move direction
if d.skill_move_type.notna().sum() == 0:
    bad("skill_move_type: never computed")
else:
    vc = d.skill_move_type.value_counts().sort_index().to_dict()
    bad_range = ((d.task_cover_best > 1.001) | (d.task_retain_best > 1.001)).any()
    if bad_range:
        bad("skill_move_type: coverage or retention exceeds 1 -- the ratio is wrong")
    elif len(vc) < 2:
        note(f"skill_move_type: only one category occurs ({vc}) -- check the 0.5 split on real data")
        ok("skill_move_type: computed and in range")
    else:
        ok(f"skill_move_type: {len(vc)} categories occur {vc}; coverage {d.task_cover_best.mean():.2f}, retention {d.task_retain_best.mean():.2f}")

# free-text occupations
for c in ("prev_occ", "target_occ"):
    s = d[c].astype(str).str.strip()
    nonblank = (s != "") & (s.str.lower() != "nan")
    if not nonblank.any():
        bad(f"{c}: no verbatim text at all")
    elif pd.to_numeric(s[nonblank], errors="coerce").notna().all():
        bad(f"{c}: looks numeric -- it should be verbatim text for NCO coding")
    else:
        ok(f"{c}: {nonblank.sum()} verbatim answers, {s[nonblank].nunique()} distinct")

# retired variables really are gone
for dead in ("location_cluster", "secondary_work_type", "secondary_work_income_pm"):
    if dead in d.columns:
        bad(f"{dead}: retired but still present in the built dataset")
if not any(x in d.columns for x in ("location_cluster", "secondary_work_type", "secondary_work_income_pm")):
    ok("retired variables (location_cluster, secondary_work_*) are gone from the built dataset")

print("\n" + "=" * 78)
print(f"{len(fails)} failure(s), {len(notes)} note(s).  Dataset: {d.shape[0]} rows x {d.shape[1]} vars")
for f in fails:
    print("  FAIL ", f)
print("=" * 78)
sys.exit(1 if fails else 0)
